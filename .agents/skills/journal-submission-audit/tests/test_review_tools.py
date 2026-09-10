"""Synthetic regression tests, not a benchmark of scientific review quality."""
import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import pymupdf as fitz
import yaml
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree as E

BASE = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('review_tools', BASE / 'scripts/review_tools.py')
tools = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tools)


def issue(fid='MAIN', quote='Test sentence.', **position):
    return dict(id='JSA-0001', priority='P2', confidence='confirmed', category='test',
                perspective=['technical'], problem='Synthetic problem', evidence='Synthetic evidence',
                impact='Synthetic impact', action='Verify the synthetic statement.', effort='verification',
                requires_author_confirmation=True, meaning_change=False,
                suggested_english='A verified test sentence.',
                locations=[dict(file_id=fid, quote=quote, **position)])


class ReviewToolsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def register(self, path, finding):
        data = dict(version='1.0', files=[dict(id='MAIN', path=path.name, role='main', sha256=tools.digest(path))], issues=[finding])
        register = self.root / 'issues.json'
        register.write_text(json.dumps(data), encoding='utf-8')
        return data, register

    def text(self, value='Test sentence.\n', name='main.txt'):
        p = self.root / name
        p.write_text(value, encoding='utf-8')
        return p

    def word(self, existing=False):
        d = Document()
        p = d.add_paragraph()
        p.add_run('Test ').bold = True
        p.add_run('sentence.')
        d.add_table(rows=1, cols=1).cell(0, 0).text = 'Table text.'
        d.sections[0].header.paragraphs[0].text = 'Unchanged header'
        if existing:
            d.add_comment(p.runs, text='Prior author note', author='Author')
        path = self.root / 'main.docx'
        d.save(path)
        return path

    def pdf(self, repeated=False, old_comment=False):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), 'Test sentence.')
        if repeated:
            page.insert_text((72, 110), 'Test sentence.')
        if old_comment:
            page.add_text_annot((280, 72), 'Prior note')
        path = self.root / 'main.pdf'
        doc.save(path)
        doc.close()
        return path

    def test_skill_metadata_and_resource_links(self):
        text = (BASE / 'SKILL.md').read_text()
        meta = yaml.safe_load(text.split('---', 2)[1])
        self.assertEqual(meta['name'], BASE.name)
        self.assertLessEqual(len(meta['description']), 1024)
        self.assertLessEqual(len(meta['compatibility']), 500)
        self.assertLess(len(text.splitlines()), 500)
        import re
        for ref in re.findall(r'\]\(([^)]+)\)', text):
            self.assertTrue((BASE / ref).is_file(), ref)

    def test_example_schema_and_all_anchors(self):
        data = tools.load(BASE / 'assets/issues.example.json', BASE)
        result = tools.run_annotations(data, BASE, None)
        self.assertEqual(result['unresolved'], 0)
        self.assertEqual(len(result['locations']), 3)

    def test_inventory_marks_extraction_not_review(self):
        p = self.text()
        result = tools.inventory(self.root, [p.name])
        self.assertEqual(result['extraction']['F001']['status'], 'extracted_not_reviewed')
        self.assertEqual(result['extraction']['F001']['lines'][0]['line'], 1)

    def test_word_inventory_includes_table_not_header(self):
        p = self.word()
        result = tools.inventory(self.root, [p.name])
        texts = [x['text'] for x in result['extraction']['F001']['paragraphs']]
        self.assertEqual(texts, ['Test sentence.', 'Table text.'])

    def test_path_traversal_rejected(self):
        with self.assertRaises(ValueError):
            tools.source(self.root, '../secret.txt')

    def test_external_symlink_rejected(self):
        with tempfile.TemporaryDirectory() as outside:
            target = Path(outside) / 'external.txt'
            target.write_text('private')
            (self.root / 'link.txt').symlink_to(target)
            with self.assertRaises(ValueError):
                tools.source(self.root, 'link.txt')

    def test_changed_hash_rejected(self):
        p = self.text()
        _, reg = self.register(p, issue(line=1))
        p.write_text('Changed manuscript.')
        with self.assertRaisesRegex(ValueError, 'hash changed'):
            tools.load(reg, self.root)

    def test_duplicate_ids_rejected(self):
        p = self.text()
        data, reg = self.register(p, issue(line=1))
        data['issues'].append(copy.deepcopy(data['issues'][0]))
        reg.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'unique'):
            tools.load(reg, self.root)

    def test_cycle_rejected(self):
        p = self.text()
        data, reg = self.register(p, issue(line=1))
        data['issues'][0]['dependencies'] = ['JSA-0001']
        reg.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'cyclic'):
            tools.load(reg, self.root)

    def test_unknown_file_rejected(self):
        p = self.text()
        _, reg = self.register(p, issue(fid='MISSING', line=1))
        with self.assertRaisesRegex(ValueError, 'unknown file'):
            tools.load(reg, self.root)

    def test_schema_diagnostic_does_not_echo_manuscript(self):
        p = self.text()
        data, reg = self.register(p, issue(line=1))
        data['issues'][0]['priority'] = 'CONFIDENTIAL EXCERPT'
        reg.write_text(json.dumps(data))
        with self.assertRaises(ValueError) as caught:
            tools.load(reg, self.root)
        self.assertNotIn('CONFIDENTIAL', str(caught.exception))

    def test_word_native_comments_preserve_original_parts(self):
        p = self.word(existing=True)
        data, _ = self.register(p, issue(paragraph=1))
        before = p.read_bytes()
        out = self.root / 'annotated'
        result = tools.run_annotations(data, self.root, out)
        self.assertEqual(result['unresolved'], 0)
        self.assertEqual(before, p.read_bytes())
        target = out / result['outputs'][0]['path']
        _, parts, document = tools.package(target)
        comments = tools.xml(parts['word/comments.xml'])
        self.assertEqual(len(comments), 2)
        self.assertIn('Prior author note', ''.join(comments.itertext()))
        self.assertIn('JSA-0001', ''.join(comments.itertext()))
        self.assertEqual(document.xpath('count(//w:commentRangeStart)', namespaces=tools.NS), 2)
        self.assertEqual([tools.ptext(x) for x in tools.paragraphs(document)], ['Test sentence.', 'Table text.'])
        with zipfile.ZipFile(p) as original, zipfile.ZipFile(target) as new:
            allowed = {'word/document.xml', 'word/comments.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
            for name in original.namelist():
                if name not in allowed:
                    self.assertEqual(original.read(name), new.read(name), name)

    def test_word_alternative_comments_part(self):
        p = self.word(existing=True)
        with zipfile.ZipFile(p) as z:
            content = {x: z.read(x) for x in z.namelist()}
        content['word/notes-comments.xml'] = content.pop('word/comments.xml')
        content['word/_rels/document.xml.rels'] = content['word/_rels/document.xml.rels'].replace(b'comments.xml', b'notes-comments.xml')
        content['[Content_Types].xml'] = content['[Content_Types].xml'].replace(b'/word/comments.xml', b'/word/notes-comments.xml')
        with zipfile.ZipFile(p, 'w') as z:
            for name, value in content.items():
                z.writestr(name, value)
        data, _ = self.register(p, issue(paragraph=1))
        out = self.root / 'alt'
        result = tools.run_annotations(data, self.root, out)
        self.assertEqual(result['unresolved'], 0)
        with zipfile.ZipFile(out / result['outputs'][0]['path']) as z:
            self.assertIn(b'JSA-0001', z.read('word/notes-comments.xml'))
            self.assertNotIn('word/comments.xml', z.namelist())

    def test_word_tracked_target_is_unresolved(self):
        p = self.word()
        d = Document(p)
        paragraph = d.paragraphs[0]
        ins = OxmlElement('w:ins')
        run = paragraph._p[-1]
        paragraph._p.remove(run)
        ins.append(run)
        paragraph._p.append(ins)
        d.save(p)
        data, _ = self.register(p, issue(paragraph=1))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_word_repeated_quote_not_guessed(self):
        p = self.word()
        d = Document(p)
        d.paragraphs[0].add_run(' Test sentence.')
        d.save(p)
        data, _ = self.register(p, issue(paragraph=1))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_pdf_native_comment_and_existing_annotation(self):
        p = self.pdf(old_comment=True)
        before = p.read_bytes()
        data, _ = self.register(p, issue(page=1))
        out = self.root / 'pdf-out'
        result = tools.run_annotations(data, self.root, out)
        self.assertEqual(result['unresolved'], 0)
        self.assertEqual(before, p.read_bytes())
        with fitz.open(out / result['outputs'][0]['path']) as d:
            self.assertEqual(len(list(d[0].annots())), 2)
            self.assertIn('Test sentence.', d[0].get_text())
            self.assertIn('JSA-0001', ' '.join(a.info['content'] for a in d[0].annots()))

    def test_pdf_repeated_quote_not_guessed(self):
        p = self.pdf(repeated=True)
        data, _ = self.register(p, issue(page=1))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_pdf_multiline_quote(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), 'First line\nsecond line.')
        p = self.root / 'main.pdf'
        doc.save(p)
        doc.close()
        data, _ = self.register(p, issue(page=1, quote='First line second line.'))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 0)

    def test_pdf_visual_box(self):
        p = self.pdf()
        data, _ = self.register(p, issue(page=1, quote='', bbox=[60, 60, 200, 120], verified_visual=True))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 0)
        data['issues'][0]['locations'][0]['bbox'] = [-50, -50, 200, 120]
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_pdf_invalid_page(self):
        p = self.pdf()
        data, _ = self.register(p, issue(page=9))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_text_bom_and_newlines_preserved(self):
        p = self.root / 'main.tex'
        p.write_bytes(b'\xef\xbb\xbfTest sentence.\r\nSecond line.\r\n')
        data, _ = self.register(p, issue(line=1))
        out = self.root / 'text-out'
        result = tools.run_annotations(data, self.root, out)
        raw = (out / result['outputs'][0]['path']).read_bytes()
        self.assertTrue(raw.startswith(b'\xef\xbb\xbf% '))
        remaining = b'\r\n'.join(line for line in raw[3:].split(b'\r\n') if not line.startswith(b'% '))
        self.assertEqual(remaining, p.read_bytes()[3:])

    def test_markdown_literal_regions(self):
        cases = [('```\nTest sentence.\n```\n', 2), ('---\nTest sentence.\n---\n', 2), ('| Test sentence. |\n', 1)]
        for text, line in cases:
            p = self.text(text, 'main.md')
            data, _ = self.register(p, issue(line=line))
            self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_tex_verbatim_region(self):
        p = self.text('\\begin{verbatim}\nTest sentence.\n\\end{verbatim}\n', 'main.tex')
        data, _ = self.register(p, issue(line=2))
        self.assertEqual(tools.run_annotations(data, self.root, None)['unresolved'], 1)

    def test_existing_outputs_refused(self):
        p = self.text()
        data, _ = self.register(p, issue(line=1))
        out = self.root / 'existing'
        out.mkdir()
        with self.assertRaises(FileExistsError):
            tools.run_annotations(data, self.root, out)
        with self.assertRaises(FileExistsError):
            tools.save(p, b'overwritten')

    def test_dtd_rejected(self):
        with self.assertRaises(ValueError):
            tools.xml(b'<!DOCTYPE x [<!ENTITY a SYSTEM "file:///secret">]><x>&a;</x>')

    def test_report_has_both_cross_file_locations(self):
        data = tools.load(BASE / 'assets/issues.example.json', BASE)
        report = tools.report(data)
        self.assertEqual(report.count('[JSA-0001 | P1'), 2)
        self.assertIn('12 s', report)
        self.assertIn('24 s', report)

    def test_cli_success_and_unresolved_exit_codes(self):
        p = self.text()
        data, reg = self.register(p, issue(line=1))
        command = [sys.executable, str(BASE / 'scripts/review_tools.py'), 'validate', str(reg), '--root', str(self.root)]
        result = subprocess.run(command, capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode())
        data['issues'][0]['locations'][0]['quote'] = 'Absent text'
        reg.write_text(json.dumps(data))
        result = subprocess.run(command, capture_output=True)
        self.assertEqual(result.returncode, 2, result.stderr.decode())


if __name__ == '__main__':
    unittest.main()
